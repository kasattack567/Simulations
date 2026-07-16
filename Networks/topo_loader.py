"""
Loader for real optical topologies from TopologyBench (Matzner et al., JOCN 17,
7-27, 2024; Zenodo 13921775). Each topology is an .xlsx with:
  - Nodes_<NAME>: Node_ID, Latitude, Longitude, Location Name, Country
  - Edges_<NAME>: Edge_ID, Source, Destination, Computed Length (km)

These are CORE optical networks (national/continental), with link lengths of
hundreds to thousands of km. We support two scales:
  - real scale : true link lengths (backbone study; CV needs many relays)
  - metro-rescaled : geometry preserved but the whole network scaled so its
                     span matches a metropolitan size (~40 km), for a like-for-
                     like DV-vs-CV metro comparison consistent with the rest of
                     the project.

The 6 selected topologies (PCA + k-means diverse selection, survivable subset):
DARKSTRAND, USA100, REDIRIS, LAMBDARAIL, NETRAIL, HIBERNIAUK.
"""
import os
import numpy as np
import openpyxl

# map short names -> filenames (edit if your filenames differ)
TOPOLOGY_FILES = {
    "DARKSTRAND": "TOP_25_DARKSTRAND.xlsx",
    "USA100":     "TOP_104_USA100.xlsx",
    "REDIRIS":    "TOP_87_REDIRIS.xlsx",
    "LAMBDARAIL": "TOP_60_LAMBDARAIL.xlsx",
    "NETRAIL":    "TOP_67_NETRAIL.xlsx",
    "HIBERNIAUK": "TOP_46_HIBERNIAUK.xlsx",
}

# folder holding the xlsx files (default: same folder as this module)
DATA_DIR = os.path.dirname(os.path.abspath(__file__))


def _haversine_km(lat1, lon1, lat2, lon2):
    """Great-circle distance in km between two lat/lon points."""
    R = 6371.0
    p1, p2 = np.radians(lat1), np.radians(lat2)
    dp = np.radians(lat2 - lat1)
    dl = np.radians(lon2 - lon1)
    a = np.sin(dp / 2) ** 2 + np.cos(p1) * np.cos(p2) * np.sin(dl / 2) ** 2
    return 2 * R * np.arcsin(np.sqrt(a))


def load_topology(name, data_dir=None):
    """Load a topology by short name.

    Returns dict with:
      name         : topology name
      node_ids     : list of node IDs (as in the file)
      idx          : {node_id: 0-based index}
      latlon       : (n,2) array of (lat, lon)
      xy_km        : (n,2) array of local planar coords in km (equirectangular
                     projection about the network centroid) — for plotting and
                     for straight-line distances
      edges        : list of (i, j) 0-based index pairs
      edge_len_km  : list of real link lengths (km) aligned with edges
    """
    data_dir = data_dir or DATA_DIR
    path = os.path.join(data_dir, TOPOLOGY_FILES.get(name, name))
    wb = openpyxl.load_workbook(path, data_only=True)
    nsheet = [s for s in wb.sheetnames if s.lower().startswith("nodes")][0]
    esheet = [s for s in wb.sheetnames if s.lower().startswith("edges")][0]

    nrows = list(wb[nsheet].iter_rows(values_only=True))[1:]
    node_ids, lat, lon = [], [], []
    for r in nrows:
        if r[0] is None:
            continue
        node_ids.append(r[0]); lat.append(float(r[1])); lon.append(float(r[2]))
    lat = np.array(lat); lon = np.array(lon)
    idx = {nid: k for k, nid in enumerate(node_ids)}

    # local planar projection (equirectangular about centroid) -> km
    lat0 = lat.mean()
    R = 6371.0
    x = np.radians(lon - lon.mean()) * R * np.cos(np.radians(lat0))
    y = np.radians(lat - lat0) * R
    xy_km = np.column_stack([x, y])

    erows = list(wb[esheet].iter_rows(values_only=True))[1:]
    edges, elen = [], []
    for r in erows:
        if r[1] is None or r[2] is None:
            continue
        s, d = r[1], r[2]
        if s in idx and d in idx:
            edges.append((idx[s], idx[d]))
            elen.append(float(r[3]) if r[3] is not None else
                        _haversine_km(lat[idx[s]], lon[idx[s]],
                                      lat[idx[d]], lon[idx[d]]))

    return dict(name=name, node_ids=node_ids, idx=idx,
                latlon=np.column_stack([lat, lon]), xy_km=xy_km,
                edges=edges, edge_len_km=elen)


def edge_lengths(topo, scale="real", metro_span_km=40.0):
    """Return an array of edge lengths under the chosen scale.

    scale='real'  : true link lengths from the file (km).
    scale='metro' : lengths rescaled so the network's maximum node separation
                    equals metro_span_km, preserving relative geometry. Uses the
                    planar xy_km coordinates (straight-line) scaled by a single
                    factor, so the whole network fits a metro footprint.
    """
    if scale == "real":
        return np.array(topo["edge_len_km"])
    # metro: scale planar coordinates so max pairwise separation = metro_span_km
    xy = topo["xy_km"]
    # network extent = max distance between any two nodes
    from itertools import combinations
    maxd = max(np.hypot(*(xy[a] - xy[b]))
               for a, b in combinations(range(len(xy)), 2))
    factor = metro_span_km / maxd if maxd > 0 else 1.0
    out = []
    for (i, j) in topo["edges"]:
        out.append(np.hypot(*(xy[i] - xy[j])) * factor)
    return np.array(out)


def scaled_xy(topo, scale="real", metro_span_km=40.0):
    """Node coordinates for plotting under the chosen scale (km)."""
    xy = topo["xy_km"] - topo["xy_km"].mean(axis=0)
    if scale == "real":
        return xy
    from itertools import combinations
    maxd = max(np.hypot(*(xy[a] - xy[b]))
               for a, b in combinations(range(len(xy)), 2))
    factor = metro_span_km / maxd if maxd > 0 else 1.0
    return xy * factor


if __name__ == "__main__":
    # quick self-check
    for name in TOPOLOGY_FILES:
        t = load_topology(name)
        real = edge_lengths(t, "real")
        metro = edge_lengths(t, "metro", 40)
        print(f"{name:12} {len(t['node_ids']):3d} nodes {len(t['edges']):3d} edges "
              f"| real {real.min():.0f}-{real.max():.0f} km "
              f"| metro {metro.min():.1f}-{metro.max():.1f} km")